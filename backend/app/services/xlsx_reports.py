"""Builds the downloadable .xlsx files from the data of `attendance_report_service`.

The layouts follow the company's own sheets (docs: "Individual Daily Report", "Daily Attendance
Report", "Leave Card"): a title line, a subtitle line, a bordered table whose Type cell is colour
coded, a total line and the legend of the type codes.

Every text that came from a person (names, remarks, reasons, a customer's name in a remark) is
written as a plain string cell - never as a formula - so a remark such as `=HYPERLINK(...)` cannot
run when the file is opened in Excel or LibreOffice.
"""

from collections.abc import Iterable
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.attendance_report_service import (
    TYPES,
    DayReport,
    EmployeeMonth,
    LeaveCard,
    MonthSummary,
    Summary,
    month_title,
)

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FONT = "Calibri"
TITLE = Font(name=FONT, size=14, bold=True)
SUBTITLE = Font(name=FONT, size=11)
HEAD = Font(name=FONT, size=11, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
BODY = Font(name=FONT, size=11)
BOLD = Font(name=FONT, size=11, bold=True)
LINK = Font(name=FONT, size=11, color="0563C1", underline="single")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
# Titles, totals and legend text sit outside the table: no wrapping, so they run across the
# (empty) cells beside them instead of being squeezed into one narrow column.
PLAIN = Alignment(horizontal="left", vertical="center", wrap_text=False)

DATE_FORMAT = "dd-mmm-yyyy ddd"
TIME_FORMAT = "hh:mm:ss"
HOURS_FORMAT = "[h]:mm"


def put_text(ws: Worksheet, row: int, col: int, value: str | None, **style) -> Cell:
    """A text cell that can never be a formula, whatever it starts with."""
    cell = ws.cell(row=row, column=col)
    cell.value = value if value else None
    if value:
        cell.data_type = "s"
    if style.get("border") is False:
        style.setdefault("align", PLAIN)
    _style(cell, **style)
    return cell


def _style(
    cell: Cell, *, font=BODY, align=LEFT, border=True, fill=None, number_format=None
) -> None:
    cell.font = font
    cell.alignment = align
    if border:
        cell.border = BORDER
    if fill:
        cell.fill = fill
    if number_format:
        cell.number_format = number_format


def _fill(code: str) -> PatternFill | None:
    colour = TYPES.get(code, ("", None))[1]
    return PatternFill("solid", fgColor=colour) if colour else None


def put_type(ws: Worksheet, row: int, col: int, code: str, label: str | None = None) -> None:
    put_text(ws, row, col, label if label is not None else code, align=CENTER, fill=_fill(code))


def put_time(ws: Worksheet, row: int, col: int, value: datetime | time | None) -> None:
    cell = ws.cell(row=row, column=col)
    if isinstance(value, datetime):
        value = value.time().replace(microsecond=0)
    cell.value = value
    _style(cell, align=CENTER, number_format=TIME_FORMAT)


def put_date(ws: Worksheet, row: int, col: int, value: date, fmt: str = DATE_FORMAT) -> None:
    cell = ws.cell(row=row, column=col)
    cell.value = value
    _style(cell, align=CENTER, number_format=fmt)


def put_number(
    ws: Worksheet, row: int, col: int, value: float | int | None, fmt: str = "General"
) -> None:
    cell = ws.cell(row=row, column=col)
    cell.value = value
    _style(cell, align=CENTER, number_format=fmt)


def put_minutes(ws: Worksheet, row: int, col: int, minutes: int | None) -> None:
    """Hours worked as an Excel duration (so a column can be summed in Excel)."""
    cell = ws.cell(row=row, column=col)
    cell.value = (minutes / 1440) if minutes is not None else None
    _style(cell, align=CENTER, number_format=HOURS_FORMAT)


def put_location(ws: Worksheet, row: int, col: int, location: str | None) -> None:
    """`lat,lng` as readable coordinates that open the place in Maps."""
    cell = ws.cell(row=row, column=col)
    if not location:
        _style(cell, align=CENTER)
        return
    try:
        lat_text, lng_text = location.split(",")
        lat, lng = float(lat_text), float(lng_text)
    except ValueError:
        put_text(ws, row, col, location, align=CENTER)
        return
    cell.value = f"{lat:.5f}, {lng:.5f}"
    cell.hyperlink = f"https://www.google.com/maps?q={lat:.6f},{lng:.6f}"
    _style(cell, font=LINK, align=CENTER)


def header_row(ws: Worksheet, row: int, titles: Iterable[str], widths: Iterable[int]) -> None:
    for col, (title, width) in enumerate(zip(titles, widths, strict=True), 1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font, cell.fill, cell.alignment, cell.border = HEAD, HEAD_FILL, CENTER, BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 32


def title_block(ws: Worksheet, title: str, subtitle: str) -> None:
    put_text(ws, 1, 1, title, font=TITLE, border=False)
    put_text(ws, 2, 1, subtitle, font=SUBTITLE, border=False)


def legend(ws: Worksheet, start_row: int, codes: Iterable[str]) -> int:
    put_text(ws, start_row, 1, "Type codes", font=BOLD, border=False)
    row = start_row + 1
    for code in codes:
        put_type(ws, row, 1, code)
        put_text(ws, row, 2, TYPES[code][0].split(" = ", 1)[1], border=False)
        row += 1
    return row


def page_setup(ws: Worksheet, header_row_number: int, landscape: bool = True) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row_number}:{header_row_number}"
    ws.freeze_panes = ws.cell(row=header_row_number + 1, column=1)


def _total_line(ws: Worksheet, row: int, text: str) -> None:
    put_text(ws, row, 1, text, font=BOLD, border=False)


def _number_text(value: float) -> str:
    return f"{value:g}"


def sheet_title(name: str, fallback: str) -> str:
    """A worksheet name from a person's name: Excel refuses \\ / ? * [ ] : in it, any name longer
    than 31 characters, an empty one, and one that starts or ends with an apostrophe."""
    cleaned = "".join("-" if char in "\\/?*[]:" or ord(char) < 32 else char for char in name)
    cleaned = cleaned.strip().strip("'").strip()[:31].strip().strip("'")
    return cleaned or fallback


def _save(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- individual month
def individual_month(report: EmployeeMonth) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title(report.user.name, "Report")
    title_block(
        ws,
        f"{report.calendar.company_name} — Individual Daily Report",
        f"Employee: {report.user.name} ({report.user.employee_code})   |   "
        f"Period: {month_title(report.first)}",
    )
    header_row(
        ws,
        4,
        ["Date", "Type", "In Punch Time", "Location at time of punch in", "Out Punch Time",
         "Location at time of Punch Out", "Remark", "Hours Worked"],
        [20, 12, 14, 26, 14, 26, 38, 13],
    )  # fmt: skip
    row = 5
    for item in report.rows:
        put_date(ws, row, 1, item.day)
        put_type(ws, row, 2, item.code) if item.code else put_text(ws, row, 2, "", align=CENTER)
        put_time(ws, row, 3, item.punch_in)
        put_location(ws, row, 4, item.location_in)
        put_time(ws, row, 5, item.punch_out)
        put_location(ws, row, 6, item.location_out)
        put_text(ws, row, 7, item.remark)
        put_minutes(ws, row, 8, item.minutes)
        row += 1
    row += 1
    _total_line(ws, row, f"Total Working Days: {_number_text(report.summary.days_worked)}")
    row = _summary_block(ws, row + 2, report.summary)
    legend(ws, row + 1, [c for c in TYPES])
    page_setup(ws, 4)
    return _save(wb)


def _summary_block(ws: Worksheet, start_row: int, s: Summary) -> int:
    put_text(ws, start_row, 1, "Summary", font=BOLD, border=False)
    items = [
        ("At Office", s.at_office), ("WFH", s.wfh), ("OD", s.od), ("WOW", s.wow), ("HW", s.hw),
        ("FD Leave", s.fd_leave), ("HD Leave", s.hd_leave), ("C/Off", s.comp_off),
        ("W/Off", s.weekly_off), ("Holiday", s.holiday), ("Absent", s.absent),
    ]  # fmt: skip
    row = start_row + 1
    for code, count in items:
        put_type(ws, row, 1, code)
        put_number(ws, row, 2, count)
        row += 1
    put_text(ws, row, 1, "Leave days taken", font=BOLD)
    put_number(ws, row, 2, s.leave_days)
    row += 1
    put_text(ws, row, 1, "Total hours worked", font=BOLD)
    put_minutes(ws, row, 2, s.total_minutes)
    row += 1
    put_text(ws, row, 1, "Average hours per day", font=BOLD)
    put_minutes(ws, row, 2, s.average_minutes)
    row += 1
    put_text(ws, row, 1, "Late arrivals", font=BOLD)
    put_number(ws, row, 2, s.late_count)
    return row + 1


# -------------------------------------------------------------------------------- day report
def day_report(report: DayReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = report.day.strftime("%d %b %Y")
    title_block(
        ws,
        f"{report.calendar.company_name} — Daily Attendance Report",
        f"Date: {report.day.strftime('%d %b %Y')} ({report.day.strftime('%A')})",
    )
    header_row(
        ws,
        4,
        ["No.", "Name", "Type", "In Punch Time", "Location at time of punch in",
         "Out Punch Time", "Location at time of Punch Out", "Remark", "Hours Worked"],
        [11, 30, 12, 14, 26, 14, 26, 38, 13],
    )  # fmt: skip
    row = 5
    for number, (user, item) in enumerate(report.people, 1):
        put_number(ws, row, 1, number)
        put_text(ws, row, 2, user.name)
        if item.code:
            put_type(ws, row, 3, item.code)
        else:
            blank = "Not in yet" if report.day == report.calendar.today else ""
            put_text(ws, row, 3, blank, align=CENTER)
        put_time(ws, row, 4, item.punch_in)
        put_location(ws, row, 5, item.location_in)
        put_time(ws, row, 6, item.punch_out)
        put_location(ws, row, 7, item.location_out)
        put_text(ws, row, 8, item.remark)
        put_minutes(ws, row, 9, item.minutes)
        row += 1
    row += 1
    _total_line(ws, row, f"Total Employees: {len(report.people)}")
    row += 2
    put_text(ws, row, 1, "Summary", font=BOLD, border=False)
    row += 1
    for code in TYPES:
        count = report.summary.counts.get(code, 0)
        if count:
            put_type(ws, row, 1, code)
            put_number(ws, row, 2, count)
            row += 1
    legend(ws, row + 1, [c for c in TYPES])
    page_setup(ws, 4)
    return _save(wb)


# --------------------------------------------------------------------------------- leave card
def leave_card(card: LeaveCard) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title(card.user.name, "Leave card")
    title_block(
        ws,
        f"{card.company_name} — Leave Card",
        f"Employee: {card.user.name} ({card.user.employee_code})   |   "
        f"Period: {card.start.strftime('%d-%b-%Y')} to {card.end.strftime('%d-%b-%Y')}",
    )
    header_row(
        ws, 4, ["Date", "Type", "Days", "In Punch", "Out Punch", "Reason"], [20, 12, 8, 13, 13, 50]
    )
    row = 5
    if not card.rows:
        put_text(ws, row, 1, "No leave records", border=False)
        row += 1
    for item in card.rows:
        put_date(ws, row, 1, item.day)
        put_type(ws, row, 2, item.code)
        put_number(ws, row, 3, item.days)
        put_time(ws, row, 4, item.punch_in)
        put_time(ws, row, 5, item.punch_out)
        put_text(ws, row, 6, item.reason)
        row += 1
    row += 1
    _total_line(ws, row, f"Total Leave Days: {_number_text(card.total_days)}")
    legend(ws, row + 2, ["FD Leave", "HD Leave", "C/Off"])
    page_setup(ws, 4, landscape=False)
    return _save(wb)


# --------------------------------------------------------------------- month summary + muster
def month_summary(report: MonthSummary) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    title_block(
        ws,
        f"{report.calendar.company_name} — Monthly Attendance Summary",
        f"Period: {month_title(report.first)}",
    )
    header_row(
        ws,
        4,
        ["No.", "Name", "Days Worked", "At Office", "WFH", "OD", "WOW", "HW", "FD Leave",
         "HD Leave",
         "C/Off", "W/Off", "Holidays", "Absent", "Total Hours", "Avg Hours / Day", "Late Arrivals",
         "Avg In-Time"],
        [11, 30, 12, 10, 8, 8, 8, 8, 10, 10, 8, 8, 10, 9, 13, 14, 12, 12],
    )  # fmt: skip
    row = 5
    for number, (user, s, _rows) in enumerate(report.people, 1):
        put_number(ws, row, 1, number)
        put_text(ws, row, 2, user.name)
        for col, value in enumerate(
            [s.days_worked, s.at_office, s.wfh, s.od, s.wow, s.hw, s.fd_leave, s.hd_leave,
             s.comp_off, s.weekly_off, s.holiday, s.absent], 3,
        ):  # fmt: skip
            put_number(ws, row, col, value)
        put_minutes(ws, row, 15, s.total_minutes)
        put_minutes(ws, row, 16, s.average_minutes)
        put_number(ws, row, 17, s.late_count)
        put_minutes(ws, row, 18, s.average_in_minute)
        row += 1
    row += 1
    _total_line(ws, row, f"Total Employees: {len(report.people)}")
    legend(ws, row + 2, [c for c in TYPES])
    page_setup(ws, 4)

    muster = wb.create_sheet("Day by day")
    title_block(
        muster,
        f"{report.calendar.company_name} — Attendance Register",
        f"Period: {month_title(report.first)}",
    )
    days = [d["date"] for d in report.daily]
    header_row(
        muster, 4, ["No.", "Name"] + [f"{d.day:02d}\n{d.strftime('%a')}" for d in days],
        [11, 30] + [9] * len(days),
    )  # fmt: skip
    for number, (user, _s, rows) in enumerate(report.people, 1):
        r = 4 + number
        put_number(muster, r, 1, number)
        put_text(muster, r, 2, user.name)
        for col, item in enumerate(rows, 3):
            cell = put_text(muster, r, col, item.code, align=CENTER, fill=_fill(item.code))
            cell.font = Font(name=FONT, size=9)
    legend(muster, 4 + len(report.people) + 2, [c for c in TYPES])
    page_setup(muster, 4)
    return _save(wb)
