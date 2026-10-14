"""Work mode at punch-in, leave kinds, the company calendar, and the attendance / leave reports
(xlsx downloads and analytics). See docs/api-contract.md sections 8, 6, 18a and 18b.

Requires a real database with migrations applied - see tests/conftest.py.
"""

from datetime import UTC, date, datetime, time, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.db.session import set_rls_context
from app.models.attendance import AttendanceDay
from app.models.leave import LeaveRequest
from app.schemas.leave import leave_days
from app.services.attendance_report_service import CompanyCalendar, classify_day
from tests.conftest import login, make_user, requires_db

IST = ZoneInfo("Asia/Kolkata")
MONTH = "2026-03"  # long past, so every day without a record is a real absence
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def ist(day: int, hour: int, minute: int, month: int = 3) -> datetime:
    return datetime(2026, month, day, hour, minute, tzinfo=IST).astimezone(UTC)


def _freeze(monkeypatch, instant: datetime):
    monkeypatch.setattr("app.services.attendance_service.utc_now", lambda: instant)


# ---------------------------------------------------------------- pure rule (no database)
def _calendar(**over) -> CompanyCalendar:
    base = {
        "company_name": "Acme",
        "weekly_off": frozenset({5, 6}),
        "holidays": {date(2026, 3, 4): "Holi"},
        "late_after": time(9, 30),
        "today": date(2026, 10, 13),
    }
    return CompanyCalendar(**{**base, **over})


class _Att:
    def __init__(
        self, mode="at_office", pin=(8, 50), pout=(18, 10), remark=None, loc="12.97,77.59"
    ):
        self.work_mode = mode
        self.remark = remark
        self.punch_in_time = ist(2, *pin) if pin else None
        self.punch_out_time = ist(2, *pout) if pout else None
        self.total_duration_minutes = 560 if pin and pout else None
        self.punch_in_location = loc
        self.punch_out_location = None


class _Leave:
    def __init__(self, kind, reason="x"):
        self.leave_kind = kind
        self.reason = reason


@pytest.mark.parametrize(
    ("day", "att", "leave", "code"),
    [
        (date(2026, 3, 2), _Att("at_office"), None, "At Office"),
        (date(2026, 3, 2), _Att("work_from_home"), None, "WFH"),
        (date(2026, 3, 2), _Att("on_duty", remark="KSB"), None, "OD"),
        (date(2026, 3, 2), _Att("at_office", pin=None, pout=None), None, "Absent"),  # no punch-in
        (date(2026, 3, 7), _Att(), None, "WOW"),  # a Saturday
        (date(2026, 3, 4), _Att(), None, "HW"),  # the holiday
        (date(2026, 3, 8), None, None, "W/Off"),  # a Sunday
        (date(2026, 3, 4), None, None, "Holiday"),
        (date(2026, 3, 5), None, None, "Absent"),
        (date(2026, 3, 5), None, _Leave("full_day"), "FD Leave"),
        (date(2026, 3, 5), None, _Leave("half_day"), "HD Leave"),
        (date(2026, 3, 5), None, _Leave("comp_off"), "C/Off"),
        (
            date(2026, 3, 5),
            _Att(),
            _Leave("full_day"),
            "FD Leave",
        ),  # leave wins, punches still shown
        (date(2026, 10, 13), None, None, ""),  # today, not punched in yet
        (date(2026, 10, 20), None, None, ""),  # the future
        (date(2026, 10, 20), None, _Leave("full_day"), "FD Leave"),  # approved leave in advance
        (date(2026, 10, 24), None, None, "W/Off"),  # weekly offs are known in advance
    ],
)
def test_one_person_day_gets_one_code(day, att, leave, code):
    if att is not None and att.punch_in_time is None:
        att = None
    row = classify_day(
        day, joined=date(2026, 1, 1), attendance=att, leave=leave, calendar=_calendar()
    )
    assert row.code == code


def test_before_the_person_joined_the_day_is_blank_unless_something_was_recorded():
    cal = _calendar()
    assert (
        classify_day(
            date(2026, 3, 5), joined=date(2026, 6, 1), attendance=None, leave=None, calendar=cal
        ).code
        == ""
    )
    assert (
        classify_day(
            date(2026, 3, 5), joined=date(2026, 6, 1), attendance=_Att(), leave=None, calendar=cal
        ).code
        == "At Office"
    )


def test_credit_hours_late_and_remarks():
    cal = _calendar()
    office = classify_day(
        date(2026, 3, 2), joined=None, attendance=_Att(), leave=None, calendar=cal
    )
    assert (office.worked, office.minutes, office.late) == (1.0, 560, False)
    late = classify_day(
        date(2026, 3, 2), joined=None, attendance=_Att(pin=(9, 45)), leave=None, calendar=cal
    )
    assert late.late is True
    on_time = classify_day(
        date(2026, 3, 2), joined=None, attendance=_Att(pin=(9, 30)), leave=None, calendar=cal
    )
    assert on_time.late is False  # exactly on the cut-off is not late
    half = classify_day(
        date(2026, 3, 5), joined=None, attendance=None, leave=_Leave("half_day"), calendar=cal
    )
    assert (half.worked, half.leave_days) == (0.5, 0.5)
    comp = classify_day(
        date(2026, 3, 5), joined=None, attendance=None, leave=_Leave("comp_off"), calendar=cal
    )
    assert (comp.worked, comp.leave_days) == (0.0, 0)
    hw = classify_day(date(2026, 3, 4), joined=None, attendance=_Att(), leave=None, calendar=cal)
    assert hw.remark == "Holi" and not hw.late  # a holiday remark when the person wrote none
    on_site = classify_day(
        date(2026, 3, 2),
        joined=None,
        attendance=_Att("on_duty", remark="KSB Coimbatore"),
        leave=None,
        calendar=cal,
    )
    assert on_site.remark == "KSB Coimbatore"


def test_leave_days_counting():
    assert leave_days("full_day", date(2026, 3, 30), date(2026, 4, 2)) == 4
    assert leave_days("half_day", date(2026, 3, 3), date(2026, 3, 3)) == 0.5
    assert leave_days("comp_off", date(2026, 3, 3), date(2026, 3, 4)) == 0


# ------------------------------------------------------------------------- the API pieces
@requires_db
async def test_punch_in_records_the_work_mode_and_a_tidy_remark(client, db_session, monkeypatch):
    await make_user(db_session, employee_code="RPT-E1", name="E", role="employee")
    headers = await login(client, employee_code="RPT-E1", password="InitialPass123!")
    _freeze(monkeypatch, ist(9, 9, 0))
    response = await client.post(
        "/api/v1/attendance/punch-in",
        json={"work_mode": "on_duty", "remark": "  KSB   Coimbatore  "},
        headers=headers,
    )
    assert response.status_code == 200
    assert response.json()["work_mode"] == "on_duty"
    assert response.json()["remark"] == "KSB Coimbatore"
    listed = await client.get("/api/v1/attendance")
    assert listed.json()["items"][0]["work_mode"] == "on_duty"


@requires_db
async def test_an_older_client_without_a_work_mode_is_at_office_and_nonsense_is_refused(
    client, db_session, monkeypatch
):
    await make_user(db_session, employee_code="RPT-E2", name="E", role="employee")
    headers = await login(client, employee_code="RPT-E2", password="InitialPass123!")
    bad = await client.post(
        "/api/v1/attendance/punch-in", json={"work_mode": "teleport"}, headers=headers
    )
    assert bad.status_code == 422
    _freeze(monkeypatch, ist(9, 9, 0))
    ok = await client.post("/api/v1/attendance/punch-in", json={}, headers=headers)
    assert ok.json()["work_mode"] == "at_office" and ok.json()["remark"] is None


@requires_db
async def test_leave_kinds_count_their_days_and_a_half_day_is_one_date(client, db_session):
    await make_user(db_session, employee_code="RPT-E3", name="E", role="employee")
    headers = await login(client, employee_code="RPT-E3", password="InitialPass123!")

    async def apply(**body):
        return await client.post("/api/v1/leaves", json=body, headers=headers)

    full = await apply(start_date="2026-12-01", end_date="2026-12-03")
    assert full.status_code == 201
    assert (full.json()["leave_kind"], full.json()["days"]) == ("full_day", 3)
    half = await apply(start_date="2026-12-07", end_date="2026-12-07", leave_kind="half_day")
    assert (half.json()["leave_kind"], half.json()["days"]) == ("half_day", 0.5)
    comp = await apply(
        start_date="2026-12-08",
        end_date="2026-12-09",
        leave_kind="comp_off",
        reason="Against 04 Dec",
    )
    assert (comp.json()["leave_kind"], comp.json()["days"]) == ("comp_off", 0)
    assert (
        await apply(start_date="2026-12-07", end_date="2026-12-08", leave_kind="half_day")
    ).status_code == 422
    assert (
        await apply(start_date="2026-12-07", end_date="2026-12-07", leave_kind="long_weekend")
    ).status_code == 422


# ---------------------------------------------------------------- company calendar endpoints
@requires_db
async def test_the_company_calendar_is_readable_by_all_and_changed_only_by_the_owner(
    client, db_session
):
    await make_user(db_session, employee_code="CAL-OWN", name="Owner", role="owner")
    await make_user(db_session, employee_code="CAL-EMP", name="Emp", role="employee")
    emp = await login(client, employee_code="CAL-EMP", password="InitialPass123!")
    seen = await client.get("/api/v1/company/settings")
    assert seen.status_code == 200
    assert seen.json()["weekly_off_days"] == [5, 6] and seen.json()["late_after"] == "09:30:00"
    assert (
        await client.patch("/api/v1/company/settings", json={"company_name": "Hacked"}, headers=emp)
    ).status_code == 403
    assert (
        await client.post(
            "/api/v1/company/holidays",
            json={"holiday_date": "2026-01-26", "name": "Republic Day"},
            headers=emp,
        )
    ).status_code == 403

    owner = await login(client, employee_code="CAL-OWN", password="InitialPass123!")
    changed = await client.patch(
        "/api/v1/company/settings",
        json={
            "company_name": "Laser Automation",
            "weekly_off_days": [6, 5, 6],
            "late_after": "09:15",
        },
        headers=owner,
    )
    assert changed.status_code == 200 and changed.json()["weekly_off_days"] == [5, 6]
    assert changed.json()["late_after"] == "09:15:00"
    assert (
        await client.patch("/api/v1/company/settings", json={"weekly_off_days": [7]}, headers=owner)
    ).status_code == 422

    added = await client.post(
        "/api/v1/company/holidays",
        json={"holiday_date": "2026-01-26", "name": " Republic Day "},
        headers=owner,
    )
    assert added.status_code == 201 and added.json()["name"] == "Republic Day"
    assert (
        await client.post(
            "/api/v1/company/holidays",
            json={"holiday_date": "2026-01-26", "name": "Again"},
            headers=owner,
        )
    ).status_code == 409
    year = await client.get("/api/v1/company/holidays?year=2026")
    assert [h["name"] for h in year.json()] == ["Republic Day"]
    removed = await client.delete(f"/api/v1/company/holidays/{added.json()['id']}", headers=owner)
    assert removed.status_code == 204
    assert (await client.get("/api/v1/company/holidays?year=2026")).json() == []
    again = await client.post(
        "/api/v1/company/holidays",
        json={"holiday_date": "2026-01-26", "name": "Republic Day"},
        headers=owner,
    )
    assert again.status_code == 201 and again.json()["id"] == added.json()["id"]  # switched back on

    actions = set((await db_session.execute(text("SELECT action FROM audit_logs"))).scalars().all())
    assert {
        "company.settings_changed",
        "company.holiday_added",
        "company.holiday_removed",
    } <= actions
    # leave the single settings row as the other tests expect it
    await client.patch(
        "/api/v1/company/settings",
        json={"weekly_off_days": [5, 6], "late_after": "09:30"},
        headers=owner,
    )


# ------------------------------------------------------------------------------ the reports
async def _company(client, db_session):
    """owner, two branches, an 'HR' manager, and a month of March 2026 for Eli."""
    owner = await make_user(db_session, employee_code="RPT-OWN", name="Olivia Owner", role="owner")
    mgr_a = await make_user(
        db_session,
        employee_code="RPT-MA",
        name="Meera Manager",
        role="manager",
        manager_id=owner.id,
    )
    mgr_b = await make_user(
        db_session,
        employee_code="RPT-MB",
        name="Bharat Manager",
        role="manager",
        manager_id=owner.id,
    )
    hr = await make_user(
        db_session, employee_code="RPT-HR", name="Hema HR", role="manager", manager_id=owner.id
    )
    eli = await make_user(
        db_session,
        employee_code="RPT-ELI",
        name="Eli Engineer",
        role="employee",
        manager_id=mgr_a.id,
    )
    ben = await make_user(
        db_session,
        employee_code="RPT-BEN",
        name="Ben Operator",
        role="employee",
        manager_id=mgr_b.id,
    )
    await db_session.execute(text("UPDATE users SET created_at = '2026-01-01'"))
    await db_session.commit()

    await set_rls_context(db_session, user_id=None, role="system")

    def day(user, d, mode, pin, pout, remark=None, loc="12.97161,77.59461", month=3):
        minutes = (pout[0] * 60 + pout[1]) - (pin[0] * 60 + pin[1]) if pout else None
        db_session.add(
            AttendanceDay(
                user_id=user.id, work_date=date(2026, month, d),
                punch_in_time=ist(d, *pin, month=month),
                punch_out_time=ist(d, *pout, month=month) if pout else None, punch_in_location=loc,
                punch_out_location="12.9,77.5" if pout else None, total_duration_minutes=minutes,
                status="present", work_mode=mode, remark=remark,
            )
        )  # fmt: skip

    day(eli, 2, "at_office", (8, 50), (18, 10))
    day(eli, 3, "work_from_home", (9, 45), (17, 0))
    day(eli, 4, "at_office", (9, 0), (17, 30))  # the Holi holiday -> HW
    day(eli, 7, "at_office", (9, 0), (13, 0))  # a Saturday -> WOW
    day(eli, 9, "on_duty", (8, 30), (18, 0), remark="KSB Coimbatore")
    day(eli, 12, "at_office", (9, 5), (14, 0))  # the half-day leave day
    day(ben, 2, "at_office", (9, 0), (18, 0))

    def leave(user, start, end, kind, reason):
        db_session.add(
            LeaveRequest(
                employee_id=user.id, start_date=start, end_date=end, leave_kind=kind, reason=reason,
                status="approved",
            )
        )  # fmt: skip

    leave(eli, date(2026, 3, 10), date(2026, 3, 11), "full_day", "Family function")
    leave(eli, date(2026, 3, 12), date(2026, 3, 12), "half_day", "Bank work")
    leave(eli, date(2026, 3, 13), date(2026, 3, 13), "comp_off", "Against working of 07 Mar")
    leave(eli, date(2026, 12, 7), date(2026, 12, 8), "full_day", "Trip")  # later: only in the card
    await db_session.commit()
    await db_session.execute(
        text(
            "INSERT INTO company_holidays (holiday_date, name) VALUES ('2026-03-04', 'Holi') ON CONFLICT DO NOTHING"  # noqa: E501
        )
    )
    await db_session.commit()
    return {"owner": owner, "mgr_a": mgr_a, "mgr_b": mgr_b, "hr": hr, "eli": eli, "ben": ben}


async def _grant_reports(client, owner_headers, hr, **extra):
    return await client.post(
        "/api/v1/permissions/grants",
        json={
            "grantee_user_id": str(hr.id),
            "permission_key": "download_company_reports",
            "scope_type": "global",
            "effect": "allow",
            **extra,
        },
        headers=owner_headers,
    )


def _sheet(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"].startswith("attachment;")
    assert response.headers["cache-control"] == "no-store"
    return load_workbook(BytesIO(response.content))


async def _statuses(client, urls):
    return {(await client.get(url)).status_code for url in urls}


def _cells(ws, row):
    return [c.value for c in ws[row]]


@requires_db
async def test_the_individual_monthly_report_has_the_companys_layout_and_the_right_codes(
    client, db_session
):
    people = await _company(client, db_session)
    await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    response = await client.get(
        f"/api/v1/reports/exports/attendance-monthly?employee_id={people['eli'].id}&month={MONTH}"
    )
    wb = _sheet(response)
    ws = wb.worksheets[0]
    assert ws["A1"].value == "Laser Automation — Individual Daily Report"
    assert ws["A2"].value == "Employee: Eli Engineer (RPT-ELI)   |   Period: March 2026"
    assert _cells(ws, 4)[:7] == [
        "Date",
        "Type",
        "In Punch Time",
        "Location at time of punch in",
        "Out Punch Time",
        "Location at time of Punch Out",
        "Remark",
    ]
    rows = {ws.cell(row=r, column=1).value.date(): [c.value for c in ws[r]] for r in range(5, 36)}
    assert len(rows) == 31
    code = {d: v[1] for d, v in rows.items()}
    assert code[date(2026, 3, 2)] == "At Office"
    assert code[date(2026, 3, 3)] == "WFH"
    assert code[date(2026, 3, 4)] == "HW" and rows[date(2026, 3, 4)][6] == "Holi"
    assert code[date(2026, 3, 7)] == "WOW"
    assert code[date(2026, 3, 9)] == "OD" and rows[date(2026, 3, 9)][6] == "KSB Coimbatore"
    assert code[date(2026, 3, 10)] == code[date(2026, 3, 11)] == "FD Leave"
    assert rows[date(2026, 3, 10)][6] == "Family function"
    assert (
        code[date(2026, 3, 12)] == "HD Leave" and rows[date(2026, 3, 12)][2] is not None
    )  # its punches are shown
    assert code[date(2026, 3, 13)] == "C/Off"
    assert code[date(2026, 3, 1)] == code[date(2026, 3, 8)] == "W/Off"
    assert code[date(2026, 3, 5)] == "Absent"
    # times are company time (08:50 IST, not 03:20 UTC) and the location links to a map
    assert rows[date(2026, 3, 2)][2] == time(8, 50)
    loc = ws.cell(row=6, column=4)
    assert loc.value == "12.97161, 77.59461" and "google.com/maps" in loc.hyperlink.target
    # the type cell is colour coded like the company's sheet
    wfh_cell = next(
        ws.cell(row=r, column=2) for r in range(5, 36) if ws.cell(row=r, column=2).value == "WFH"
    )
    assert wfh_cell.fill.fgColor.rgb.endswith("70AD47")
    text_cells = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "Total Working Days: 5.5" in text_cells  # 2,3,4,7,9 = 5, plus the half-day leave 0.5
    assert "Work From Home" in text_cells and "Type codes" in text_cells  # the legend
    # hours are an Excel duration
    assert rows[date(2026, 3, 2)][7] == timedelta(minutes=560)


@requires_db
async def test_the_daily_report_lists_everyone_with_their_type_and_a_total(client, db_session):
    await _company(client, db_session)
    await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    wb = _sheet(await client.get("/api/v1/reports/exports/attendance-daily?date=2026-03-02"))
    ws = wb.worksheets[0]
    assert ws["A1"].value == "Laser Automation — Daily Attendance Report"
    assert ws["A2"].value == "Date: 02 Mar 2026 (Monday)"
    header = _cells(ws, 4)[:8]
    assert header == [
        "No.",
        "Name",
        "Type",
        "In Punch Time",
        "Location at time of punch in",
        "Out Punch Time",
        "Location at time of Punch Out",
        "Remark",
    ]
    body = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in range(5, 12)}
    assert body["Eli Engineer"] == "At Office" and body["Ben Operator"] == "At Office"
    assert body["Meera Manager"] == "Absent" and body["Hema HR"] == "Absent"
    assert "Olivia Owner" not in body  # the Owner is not a reporting subject unless they punched
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "Total Employees: 5" in texts  # everyone but the Owner
    wb = _sheet(await client.get("/api/v1/reports/exports/attendance-daily?date=2026-03-10"))
    rows = {r[1].value: r[2].value for r in wb.worksheets[0].iter_rows(min_row=5, max_row=9)}
    assert rows["Eli Engineer"] == "FD Leave"
    wb = _sheet(await client.get("/api/v1/reports/exports/attendance-daily?date=2026-03-08"))
    assert {r[2].value for r in wb.worksheets[0].iter_rows(min_row=5, max_row=9)} == {"W/Off"}


@requires_db
async def test_the_leave_card_lists_each_leave_day_and_totals_them(client, db_session):
    people = await _company(client, db_session)
    await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    wb = _sheet(
        await client.get(
            f"/api/v1/reports/exports/leave-card?employee_id={people['eli'].id}&start_date=2026-01-01&end_date=2026-12-31"
        )
    )
    ws = wb.worksheets[0]
    assert ws["A1"].value == "Laser Automation — Leave Card"
    assert _cells(ws, 4)[:6] == ["Date", "Type", "Days", "In Punch", "Out Punch", "Reason"]
    rows = [[c.value for c in ws[r]] for r in range(5, 11)]
    assert [(r[0].date(), r[1], r[2], r[5]) for r in rows] == [
        (date(2026, 3, 10), "FD Leave", 1, "Family function"),
        (date(2026, 3, 11), "FD Leave", 1, "Family function"),
        (date(2026, 3, 12), "HD Leave", 0.5, "Bank work"),
        (date(2026, 3, 13), "C/Off", 0, "Against working of 07 Mar"),
        (date(2026, 12, 7), "FD Leave", 1, "Trip"),
        (date(2026, 12, 8), "FD Leave", 1, "Trip"),
    ]
    assert rows[2][3] == time(9, 5)  # the half day shows when they actually came in
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "Total Leave Days: 4.5" in texts  # 2 + 0.5 + 0 + 2
    # the default period is the current financial year (April-March) - an old card is empty there
    empty = _sheet(
        await client.get(f"/api/v1/reports/exports/leave-card?employee_id={people['ben'].id}")
    )
    assert any(
        c.value == "No leave records" for row in empty.worksheets[0].iter_rows() for c in row
    )
    assert (
        await client.get(
            f"/api/v1/reports/exports/leave-card?employee_id={people['eli'].id}&start_date=2026-05-01&end_date=2026-04-01"
        )
    ).status_code == 422


@requires_db
async def test_the_monthly_summary_has_a_row_per_person_and_a_day_by_day_register(
    client, db_session
):
    await _company(client, db_session)
    await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    wb = _sheet(await client.get(f"/api/v1/reports/exports/attendance-summary?month={MONTH}"))
    assert wb.sheetnames == ["Summary", "Day by day"]
    ws = wb["Summary"]
    head = _cells(ws, 4)
    assert head[:6] == ["No.", "Name", "Days Worked", "At Office", "WFH", "OD"]
    eli = next(r for r in ws.iter_rows(min_row=5, max_row=8) if r[1].value == "Eli Engineer")
    got = dict(zip(head, [c.value for c in eli], strict=False))
    assert (got["Days Worked"], got["At Office"], got["WFH"], got["OD"], got["WOW"], got["HW"]) == (
        5.5,
        1,
        1,
        1,
        1,
        1,
    )  # 2 March at the office; the 4th is HW, the 7th WOW and the 12th a half-day leave
    assert (got["FD Leave"], got["HD Leave"], got["C/Off"], got["Late Arrivals"]) == (2, 1, 1, 1)
    muster = wb["Day by day"]
    assert muster.cell(row=4, column=3).value.startswith("01")
    eli_row = next(
        r for r in muster.iter_rows(min_row=5, max_row=8) if r[1].value == "Eli Engineer"
    )
    assert eli_row[3].value == "At Office"  # 2 March is the second day column (index 3)


@requires_db
async def test_text_from_people_never_becomes_a_formula_in_the_sheet(client, db_session):
    people = await _company(client, db_session)
    await db_session.execute(
        text(
            "UPDATE users SET name = '=HYPERLINK(\"http://evil\",\"x\")' WHERE employee_code = 'RPT-ELI'"  # noqa: E501
        )
    )
    await db_session.execute(
        text("UPDATE attendance_days SET remark = '=1+1' WHERE work_date = '2026-03-09'")
    )
    await db_session.execute(
        text("UPDATE leave_requests SET reason = '@SUM(1)' WHERE reason = 'Family function'")
    )
    await db_session.commit()
    await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    for url in (
        f"/api/v1/reports/exports/attendance-monthly?employee_id={people['eli'].id}&month={MONTH}",
        "/api/v1/reports/exports/attendance-daily?date=2026-03-09",
        f"/api/v1/reports/exports/leave-card?employee_id={people['eli'].id}&start_date=2026-03-01&end_date=2026-03-31",
    ):
        wb = _sheet(await client.get(url))
        formulas = [
            c.coordinate for ws in wb for row in ws.iter_rows() for c in row if c.data_type == "f"
        ]
        assert formulas == [], (url, formulas)
    monthly = _sheet(
        await client.get(
            f"/api/v1/reports/exports/attendance-monthly?employee_id={people['eli'].id}&month={MONTH}"
        )
    ).worksheets[0]
    assert monthly["A2"].value.startswith("Employee: =HYPERLINK")
    assert any(c.value == "=1+1" for row in monthly.iter_rows() for c in row)


# ----------------------------------------------------------------------------------- access
@requires_db
async def test_only_the_owner_and_permitted_managers_may_use_the_company_reports(
    client, db_session
):
    people = await _company(client, db_session)
    urls = [
        f"/api/v1/reports/exports/attendance-monthly?employee_id={people['eli'].id}&month={MONTH}",
        "/api/v1/reports/exports/attendance-daily?date=2026-03-02",
        f"/api/v1/reports/exports/leave-card?employee_id={people['eli'].id}",
        f"/api/v1/reports/exports/attendance-summary?month={MONTH}",
        f"/api/v1/reports/attendance-analytics?month={MONTH}",
        "/api/v1/reports/leave-analytics",
    ]
    assert await _statuses(client, urls) == {401}

    await login(client, employee_code="RPT-ELI", password="InitialPass123!")
    assert await _statuses(client, urls) == {403}
    await client.post(
        "/api/v1/auth/logout", headers={"X-CSRF-Token": client.cookies.get("csrf_token")}
    )

    await login(client, employee_code="RPT-MA", password="InitialPass123!")  # a plain manager
    assert await _statuses(client, urls) == {403}
    await client.post(
        "/api/v1/auth/logout", headers={"X-CSRF-Token": client.cookies.get("csrf_token")}
    )

    owner = await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    assert await _statuses(client, urls) == {200}
    granted = await _grant_reports(client, owner, people["hr"])
    assert granted.status_code == 201
    await client.post("/api/v1/auth/logout", headers=owner)

    await login(client, employee_code="RPT-HR", password="InitialPass123!")
    assert await _statuses(client, urls) == {200}  # company-wide, any branch
    ben_card = await client.get(
        f"/api/v1/reports/exports/leave-card?employee_id={people['ben'].id}"
    )
    assert ben_card.status_code == 200
    # ... but only through the reports: the same person still sees just their own team elsewhere
    assert (await client.get(f"/api/v1/users/{people['ben'].id}")).status_code in (403, 404)
    assert (await client.get("/api/v1/attendance")).json()["total"] == 0


@requires_db
async def test_a_denial_or_a_revoke_takes_the_report_access_away(client, db_session):
    people = await _company(client, db_session)
    owner = await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    allow = (await _grant_reports(client, owner, people["hr"])).json()
    deny = (await _grant_reports(client, owner, people["hr"], effect="deny")).json()
    await client.post("/api/v1/auth/logout", headers=owner)
    url = f"/api/v1/reports/attendance-analytics?month={MONTH}"
    await login(client, employee_code="RPT-HR", password="InitialPass123!")
    assert (await client.get(url)).status_code == 403  # a deny always wins
    await client.post(
        "/api/v1/auth/logout", headers={"X-CSRF-Token": client.cookies.get("csrf_token")}
    )

    owner = await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    await client.delete(f"/api/v1/permissions/grants/{deny['id']}", headers=owner)
    await client.post("/api/v1/auth/logout", headers=owner)
    await login(client, employee_code="RPT-HR", password="InitialPass123!")
    assert (await client.get(url)).status_code == 200
    await client.post(
        "/api/v1/auth/logout", headers={"X-CSRF-Token": client.cookies.get("csrf_token")}
    )

    owner = await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    await client.delete(f"/api/v1/permissions/grants/{allow['id']}", headers=owner)
    await client.post("/api/v1/auth/logout", headers=owner)
    await login(client, employee_code="RPT-HR", password="InitialPass123!")
    assert (await client.get(url)).status_code == 403


@requires_db
async def test_every_download_is_audited_and_bad_input_is_refused(client, db_session):
    people = await _company(client, db_session)
    await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    ok = await client.get(
        f"/api/v1/reports/exports/attendance-monthly?employee_id={people['eli'].id}&month={MONTH}"
    )
    assert "Attendance_Eli_Engineer_March_2026.xlsx" in ok.headers["content-disposition"]
    await client.get("/api/v1/reports/exports/attendance-daily?date=2026-03-02")
    entries = (
        (
            await db_session.execute(
                text(
                    "SELECT new_value FROM audit_logs WHERE action = 'report.downloaded' ORDER BY created_at"  # noqa: E501
                )
            )
        )
        .scalars()
        .all()
    )
    assert [e["report"] for e in entries] == ["attendance_monthly", "attendance_daily"]
    assert entries[0]["employee_id"] == str(people["eli"].id)
    for url in (
        "/api/v1/reports/exports/attendance-summary?month=2026-13",
        "/api/v1/reports/exports/attendance-summary?month=March",
        "/api/v1/reports/exports/attendance-daily?date=nonsense",
        "/api/v1/reports/exports/attendance-monthly?employee_id=not-a-uuid&month=2026-03",
    ):
        assert (await client.get(url)).status_code == 422, url
    missing = await client.get(
        f"/api/v1/reports/exports/attendance-monthly?employee_id=00000000-0000-0000-0000-000000000000&month={MONTH}"
    )
    assert missing.status_code == 404


# ---------------------------------------------------------------------------------- analytics
@requires_db
async def test_attendance_analytics_numbers(client, db_session):
    await _company(client, db_session)
    await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    body = (await client.get(f"/api/v1/reports/attendance-analytics?month={MONTH}")).json()
    assert (body["month"], body["employees"]) == (MONTH, 5)  # everyone but the Owner
    # 22 weekdays in March 2026; the holiday (4 Mar) is not a working day -> 21
    assert body["working_days"] == 21
    eli = next(p for p in body["people"] if p["employee_code"] == "RPT-ELI")
    assert eli["days_worked"] == 5.5 and eli["work_from_home"] == 1 and eli["on_duty"] == 1
    assert (
        eli["weekend_or_holiday_working"] == 2 and eli["leave_days"] == 2.5 and eli["comp_off"] == 1
    )
    assert eli["late_count"] == 1 and eli["average_minutes"] is not None
    day2 = next(d for d in body["daily"] if d["date"] == "2026-03-02")
    assert (day2["at_office"], day2["absent"]) == (2, 3) and not day2["weekly_off"]
    assert next(d for d in body["daily"] if d["date"] == "2026-03-04")["holiday"] is True
    assert next(d for d in body["daily"] if d["date"] == "2026-03-07")["weekly_off"] is True
    assert body["late_arrivals"] == 1
    codes = {t["code"]: t["count"] for t in body["by_type"]}
    assert codes["At Office"] == 2 and codes["WFH"] == 1 and codes["FD Leave"] == 2
    assert 0 < body["average_attendance_percent"] < 20  # two of five people worked a few days


@requires_db
async def test_leave_analytics_by_financial_year(client, db_session, monkeypatch):
    monkeypatch.setattr(
        "app.services.attendance_report_service.company_today", lambda: date(2026, 10, 14)
    )
    await _company(client, db_session)
    await login(client, employee_code="RPT-OWN", password="InitialPass123!")
    fy25 = (
        await client.get("/api/v1/reports/leave-analytics?financial_year=2025")
    ).json()  # Apr 2025 - Mar 2026
    assert (fy25["first_day"], fy25["last_day"]) == ("2025-04-01", "2026-03-31")
    assert fy25["total_leave_days"] == 2.5  # 10-11 Mar full, 12 Mar half (0.5), comp-off 0
    assert [m["month"] for m in fy25["by_month"]][:2] == ["2025-04", "2025-05"] and len(
        fy25["by_month"]
    ) == 12
    march = next(m for m in fy25["by_month"] if m["month"] == "2026-03")
    assert (
        march["leave_days"],
        march["full_days"],
        march["half_days"],
        march["comp_off_days"],
    ) == (2.5, 2, 1, 1)
    fy26 = (await client.get("/api/v1/reports/leave-analytics?financial_year=2026")).json()
    assert fy26["total_leave_days"] == 2 and fy26["people"][0]["employee_code"] == "RPT-ELI"
    assert (
        fy26["upcoming_days"] == 2 and fy26["taken_days"] == 0
    )  # the December trip is still ahead


@requires_db
async def test_the_new_permission_is_in_the_catalogue(client, db_session):
    await make_user(db_session, employee_code="CAT-OWN", name="Owner", role="owner")
    owner = await login(client, employee_code="CAT-OWN", password="InitialPass123!")
    keys = {
        d["key"]
        for d in (await client.get("/api/v1/permissions/definitions", headers=owner)).json()
    }
    assert "download_company_reports" in keys
