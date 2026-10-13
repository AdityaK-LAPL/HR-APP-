"""Builds UAT-scenarios.md and UAT-scenarios.xlsx from scenarios.py.

    pip install openpyxl
    python build_uat.py
"""

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import scenarios as S

HERE = Path(__file__).parent
RESULTS = ["Pass", "Fail", "Blocked", "N/A"]


def numbered(steps):
    return [f"{i}. {s}" for i, s in enumerate(steps, 1)]


def esc(text):
    return text.replace("|", "\\|")


# ------------------------------------------------------------------ Markdown
def build_markdown():
    n = sum(len(rows) for _, _, rows in S.SECTIONS)
    out = [
        "# UAT scenarios",
        "",
        f"{n} scenarios. Generated from `scenarios.py` by `build_uat.py` - edit there, not here. "
        "Testers normally work in the spreadsheet `UAT-scenarios.xlsx` (same content, with drop-down "
        "Pass/Fail cells and a summary); this file is the printable / reviewable copy. "
        "How to run the exercise: [README.md](README.md).",
        "",
        "**Order:** run the *Owner* scenarios first (UAT-OWN-01 to 05 create the people everybody else uses), "
        "then Manager, then Employee, then the cross-cutting checks. Scenarios that need another scenario's "
        "result say so in their steps.",
        "",
        "Result values: **Pass**, **Fail**, **Blocked** (could not run), **N/A**. Write what actually happened in "
        "*Actual result* (always for a Fail), anything else in *Notes*.",
        "",
        "## The people (the \"cast\")",
        "",
        "| Employee ID | Name | Role | Reports to | Designation | Used for |",
        "|---|---|---|---|---|---|",
    ]
    out += [f"| {' | '.join(esc(c) for c in row)} |" for row in S.CAST]
    out += ["", "## Business data used by the scenarios", "", "| Item | Details |", "|---|---|"]
    out += [f"| {a} | {esc(b)} |" for a, b in S.DATA]
    for title, _role, rows in S.SECTIONS:
        out += ["", f"## {title}", "", "| Test ID | Role | Area | Steps | Expected result | Actual result | Pass / Fail | Notes |", "|---|---|---|---|---|---|---|---|"]
        for sc in rows:
            sid, role, area, name, steps, expected = sc[:6]
            device = sc[6] if len(sc) > 6 else ""
            steps_html = "<br>".join(esc(s) for s in numbered(steps))
            exp_html = "<br>".join("- " + esc(e) for e in expected)
            label = f"**{esc(name)}**" + (f"<br>*Device: {device}*" if device else "")
            out.append(f"| {sid} | {role} | {area} | {label}<br>{steps_html} | {exp_html} |  |  |  |")
    return "\n".join(out) + "\n"


# ------------------------------------------------------------------ Excel
HEAD_FILL = PatternFill("solid", fgColor="1E3A8A")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, row, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row=row, column=i, value=t)
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center"), BORDER
        ws.column_dimensions[get_column_letter(i)].width = w


def put_rows(ws, start, rows):
    for r, row in enumerate(rows, start):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment, cell.border = WRAP, BORDER


def build_xlsx(path):
    wb = Workbook()

    ws = wb.active
    ws.title = "Start here"
    lines = [
        ("UAT - company management and service app", True),
        ("", False),
        ("1. Fill your name, the device and the browser on the 'Cast & setup' sheet.", False),
        ("2. Run the Owner sheet first (UAT-OWN-01 to 05 create the people), then Manager, Employee and Cross-cutting.", False),
        ("3. For each scenario do the steps, then write what happened in 'Actual result' and choose Pass / Fail / Blocked / N/A in the 'Result' column.", False),
        ("4. For every Fail also add a line on the 'Defects' sheet (what, how to repeat it, how bad). A Fail with a clear Actual result is enough to start with.", False),
        ("5. The 'Summary' sheet counts the results by itself. Sign-off is at the bottom of 'Summary'.", False),
        ("", False),
        ("Severity: Blocker = cannot continue or data is wrong/lost or someone sees what they must not. Major = a workflow does not work but there is a workaround. Minor = works but is awkward or unclear. Cosmetic = looks wrong only.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(bold=bold, size=14 if bold else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130

    ws = wb.create_sheet("Cast & setup")
    header(ws, 1, ["Employee ID", "Name", "Role", "Reports to", "Designation", "Used for", "Temporary password (write it here, then change)"], [14, 22, 12, 18, 20, 55, 40])
    put_rows(ws, 2, [list(r) + [""] for r in S.CAST])
    r0 = len(S.CAST) + 4
    ws.cell(row=r0, column=1, value="Business data used by the scenarios").font = Font(bold=True)
    header(ws, r0 + 1, ["Item", "Details"], [14, 22])
    put_rows(ws, r0 + 2, [list(r) for r in S.DATA])
    r1 = r0 + len(S.DATA) + 4
    ws.cell(row=r1, column=1, value="Tester").font = Font(bold=True)
    for i, label in enumerate(["Tester name", "Role under test", "Device and model", "Browser and version", "Date started", "Web address used"], 1):
        ws.cell(row=r1 + i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r1 + i, column=2).border = BORDER
        ws.merge_cells(start_row=r1 + i, start_column=2, end_row=r1 + i, end_column=4)

    sheet_names = []
    for title, _role, rows in S.SECTIONS:
        name = title.split(",")[0].replace(" and other cross-cutting checks", "")
        name = "Cross-cutting" if name.startswith("Mobile") else name
        sheet_names.append(name)
        ws = wb.create_sheet(name)
        header(ws, 1, ["Test ID", "Role", "Area", "Device", "Scenario", "Steps", "Expected result", "Actual result", "Result", "Notes", "Tester", "Date"],
               [13, 10, 14, 9, 30, 62, 62, 40, 11, 30, 12, 12])
        ws.freeze_panes = "F2"
        for r, sc in enumerate(rows, 2):
            sid, role, area, nm, steps, expected = sc[:6]
            device = sc[6] if len(sc) > 6 else ""
            vals = [sid, role, area, device, nm, "\n".join(numbered(steps)), "\n".join("- " + e for e in expected), "", "", "", "", ""]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.alignment, cell.border = WRAP, BORDER
            ws.cell(row=r, column=1).font = Font(bold=True)
            ws.cell(row=r, column=5).font = Font(bold=True)
        last = len(rows) + 1
        dv = DataValidation(type="list", formula1='"' + ",".join(RESULTS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"I2:I{last}")
        rng = f"I2:I{last}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", bgColor="C6EFCE")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fail"'], fill=PatternFill("solid", bgColor="FFC7CE")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Blocked"'], fill=PatternFill("solid", bgColor="FFEB9C")))
        ws.auto_filter.ref = f"A1:L{last}"

    ws = wb.create_sheet("Defects")
    header(ws, 1, ["Defect #", "Test ID", "Severity", "What went wrong", "How to repeat it", "Found by", "Date", "Status", "Fix / decision"], [10, 13, 12, 50, 50, 14, 12, 12, 40])
    dv = DataValidation(type="list", formula1='"Blocker,Major,Minor,Cosmetic"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("C2:C200")
    dv2 = DataValidation(type="list", formula1='"Open,Fixed,Retested OK,Deferred,Rejected"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("H2:H200")
    for r in range(2, 201):
        for c in range(1, 10):
            ws.cell(row=r, column=c).border = BORDER
            ws.cell(row=r, column=c).alignment = WRAP
        ws.cell(row=r, column=1, value=r - 1)

    ws = wb.create_sheet("Summary", 1)
    header(ws, 1, ["Sheet", "Scenarios", "Pass", "Fail", "Blocked", "N/A", "Not run", "% done"], [28, 11, 9, 9, 9, 9, 9, 9])
    for i, name in enumerate(sheet_names, 2):
        n = len(S.SECTIONS[i - 2][2])
        q = f"'{name}'!I2:I{n + 1}"
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=n)
        for c, label in enumerate(RESULTS, 3):
            ws.cell(row=i, column=c, value=f'=COUNTIF({q},"{label}")')
        ws.cell(row=i, column=7, value=f"=B{i}-SUM(C{i}:F{i})")
        ws.cell(row=i, column=8, value=f"=IF(B{i}=0,0,SUM(C{i}:F{i})/B{i})").number_format = "0%"
    t = len(sheet_names) + 2
    ws.cell(row=t, column=1, value="Total").font = Font(bold=True)
    for c in range(2, 8):
        col = get_column_letter(c)
        ws.cell(row=t, column=c, value=f"=SUM({col}2:{col}{t - 1})").font = Font(bold=True)
    ws.cell(row=t, column=8, value=f"=IF(B{t}=0,0,SUM(C{t}:F{t})/B{t})").number_format = "0%"
    ws.cell(row=t + 2, column=1, value="Open blockers").font = Font(bold=True)
    ws.cell(row=t + 2, column=2, value='=COUNTIFS(Defects!C2:C200,"Blocker",Defects!H2:H200,"Open")')
    ws.cell(row=t + 3, column=1, value="Open majors").font = Font(bold=True)
    ws.cell(row=t + 3, column=2, value='=COUNTIFS(Defects!C2:C200,"Major",Defects!H2:H200,"Open")')
    ws.cell(row=t + 5, column=1, value="Exit criteria (agree before starting): no open Blocker, no open Major without an agreed workaround, every scenario run (Pass, or Fail with a defect, or N/A with a reason).").alignment = WRAP
    ws.merge_cells(start_row=t + 5, start_column=1, end_row=t + 5, end_column=8)
    ws.row_dimensions[t + 5].height = 45
    ws.cell(row=t + 7, column=1, value="Sign-off").font = Font(bold=True)
    for i, label in enumerate(["Accepted / not accepted", "Name", "Role", "Date", "Signature"], t + 8):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2).border = BORDER
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=5)
    wb.save(path)


if __name__ == "__main__":
    (HERE / "UAT-scenarios.md").write_text(build_markdown(), encoding="utf-8")
    build_xlsx(HERE / "UAT-scenarios.xlsx")
    total = sum(len(r) for _, _, r in S.SECTIONS)
    print(f"{total} scenarios written")
